import json
import logging
import os
from pathlib import Path
from typing import Callable, List, Optional, Tuple

import pdfplumber
from docx import Document as DocxDocument
from openai import OpenAI

from helpers.doc_type import infer_doc_type
from helpers.embeddings import embed_text
from helpers.jd_validator import load_combined_hr_jd_text, validate_jd_text
from helpers.text_chunker import split_text
from models.schemas import (
    CollectionName,
    DocumentChunk,
    IngestFileResult,
    IngestResponse,
)
from services.vector_store import ChromaVectorStore

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".json", ".xlsx", ".xls"}


class DocumentIngestionService:
    def __init__(self, vector_store: ChromaVectorStore, client: OpenAI, config: dict):
        self._store = vector_store
        self._client = client
        self._config = config

    def ingest_system(
        self,
        folder_path: Optional[str] = None,
        progress_callback: Optional[Callable[[str, float], None]] = None,
    ) -> IngestResponse:
        base = self._config.get("data_folder", "data_RAG")
        folder = folder_path or os.path.join(os.getcwd(), base, "system")
        return self._ingest_collection(
            folder=folder,
            collection="system",
            knowledge_base="system",
            owner_id=None,
            progress_callback=progress_callback,
            clear_message="system_kb",
        )

    def ingest_hr(
        self,
        owner_id: str,
        folder_path: Optional[str] = None,
        progress_callback: Optional[Callable[[str, float], None]] = None,
    ) -> IngestResponse:
        if not owner_id or not owner_id.strip():
            return IngestResponse(success=False, message="owner_id không hợp lệ.")
        owner_id = owner_id.strip()
        base = self._config.get("data_folder", "data_RAG")
        folder = folder_path or os.path.join(os.getcwd(), base, "hr", owner_id)
        return self._ingest_collection(
            folder=folder,
            collection="hr",
            knowledge_base="hr",
            owner_id=owner_id,
            progress_callback=progress_callback,
            clear_message=f"hr_kb (owner={owner_id})",
        )

    def ingest_uploaded_files(
        self,
        collection: CollectionName,
        uploads: List[Tuple[str, bytes]],
        owner_id: Optional[str] = None,
    ) -> IngestResponse:
        if collection == "hr":
            if not owner_id or not owner_id.strip():
                return IngestResponse(success=False, message="owner_id không hợp lệ.")
            owner_id = owner_id.strip()

        if not uploads:
            return IngestResponse(success=False, message="Cần ít nhất một file.")

        max_files = self._config.get("max_upload_files", 20)
        if len(uploads) > max_files:
            return IngestResponse(
                success=False,
                message=f"Vượt quá giới hạn {max_files} file mỗi request.",
            )

        knowledge_base = collection
        clear_message = "system_kb" if collection == "system" else f"hr_kb (owner={owner_id})"
        target_dir = self._upload_target_dir(collection, owner_id)
        os.makedirs(target_dir, exist_ok=True)

        file_results: List[IngestFileResult] = []
        succeeded_files: List[str] = []
        failed_files: List[str] = []
        total_chunks = 0

        for file_name, content in uploads:
            safe_name = self._sanitize_filename(file_name)
            if not safe_name:
                result = IngestFileResult(
                    file_name=file_name or "(unknown)",
                    success=False,
                    message="Tên file không hợp lệ.",
                )
                file_results.append(result)
                failed_files.append(file_name or "(unknown)")
                continue

            ext_error = self._validate_extension(safe_name)
            if ext_error:
                result = IngestFileResult(
                    file_name=safe_name,
                    success=False,
                    message=ext_error,
                )
                file_results.append(result)
                failed_files.append(safe_name)
                continue

            size_error = self._validate_file_size(len(content))
            if size_error:
                result = IngestFileResult(
                    file_name=safe_name,
                    success=False,
                    message=size_error,
                )
                file_results.append(result)
                failed_files.append(safe_name)
                continue

            dest_path = os.path.join(target_dir, safe_name)
            try:
                Path(dest_path).write_bytes(content)
                if collection == "hr":
                    raw_text = self._extract_text(dest_path)
                    validation = validate_jd_text(raw_text, safe_name, self._config)
                    if not validation.valid:
                        Path(dest_path).unlink(missing_ok=True)
                        msg = "; ".join(validation.errors)
                        result = IngestFileResult(
                            file_name=safe_name,
                            success=False,
                            chunks_created=0,
                            message=msg,
                            validation_errors=validation.errors,
                        )
                        file_results.append(result)
                        failed_files.append(safe_name)
                        continue
                chunks_created = self._ingest_single_file_path(
                    file_path=dest_path,
                    collection=collection,
                    knowledge_base=knowledge_base,
                    owner_id=owner_id,
                    progress_callback=None,
                    file_index=1,
                    total_files=1,
                )
                if chunks_created == 0:
                    result = IngestFileResult(
                        file_name=safe_name,
                        success=False,
                        chunks_created=0,
                        message="Không trích xuất được nội dung hoặc không tạo chunk.",
                    )
                    failed_files.append(safe_name)
                else:
                    result = IngestFileResult(
                        file_name=safe_name,
                        success=True,
                        chunks_created=chunks_created,
                        message="OK",
                    )
                    succeeded_files.append(safe_name)
                    total_chunks += chunks_created
                file_results.append(result)
            except Exception as e:
                logger.exception("Upload ingest error %s", safe_name)
                result = IngestFileResult(
                    file_name=safe_name,
                    success=False,
                    message=str(e),
                )
                file_results.append(result)
                failed_files.append(safe_name)

        loaded = len(succeeded_files)
        if loaded == 0:
            return IngestResponse(
                success=False,
                message=f"Không index được file nào ({clear_message}).",
                documents_loaded=0,
                chunks_created=0,
                files=[],
                failed_files=failed_files,
                file_results=file_results,
            )

        total = len(uploads)
        return IngestResponse(
            success=True,
            message=f"Index {loaded}/{total} file, {total_chunks} chunks ({clear_message}).",
            documents_loaded=loaded,
            chunks_created=total_chunks,
            files=succeeded_files,
            failed_files=failed_files,
            file_results=file_results,
        )

    def load_combined_hr_jd_text(self, owner_id: str) -> str:
        base = self._config.get("data_folder", "data_RAG")
        return load_combined_hr_jd_text(base, owner_id, self._extract_text)

    def check_existing_data(self) -> int:
        system_count = self._store.count("system")
        hr_count = self._store.count("hr")
        total = system_count + hr_count
        if total > 0:
            print(f"ChromaDB: system_kb={system_count} chunks, hr_kb={hr_count} chunks.")
        else:
            print("Chưa có dữ liệu. Dùng 'ingest-system' hoặc 'ingest-hr <user_id>'.")
        return total

    def _upload_target_dir(self, collection: CollectionName, owner_id: Optional[str]) -> str:
        base = self._config.get("data_folder", "data_RAG")
        root = os.path.join(os.getcwd(), base)
        if collection == "system":
            return os.path.join(root, "system")
        return os.path.join(root, "hr", owner_id or "")

    def _sanitize_filename(self, name: str) -> str:
        if not name or not name.strip():
            return ""
        safe = Path(name.strip()).name
        if not safe or safe in (".", "..") or ".." in safe:
            return ""
        return safe

    def _validate_extension(self, file_name: str) -> Optional[str]:
        ext = Path(file_name).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            return f"Định dạng '{ext}' không hỗ trợ. Hỗ trợ: {supported}"
        return None

    def _validate_file_size(self, size_bytes: int) -> Optional[str]:
        max_mb = self._config.get("max_upload_size_mb", 25)
        max_bytes = max_mb * 1024 * 1024
        if size_bytes > max_bytes:
            return f"File vượt quá {max_mb} MB."
        return None

    def _ingest_collection(
        self,
        folder: str,
        collection: CollectionName,
        knowledge_base: str,
        owner_id: Optional[str],
        progress_callback: Optional[Callable[[str, float], None]],
        clear_message: str,
    ) -> IngestResponse:
        if not os.path.exists(folder):
            os.makedirs(folder, exist_ok=True)
            return IngestResponse(
                success=False,
                message=f"Đã tạo folder '{folder}'. Cho file vào và gọi lại.",
            )

        self._emit_progress(progress_callback, f"Xóa dữ liệu cũ trong {clear_message}...", 0.0)
        deleted = self._store.delete_scoped(collection, owner_id=owner_id)
        self._emit_progress(progress_callback, f"Đã xóa {deleted} bản ghi cũ.", 0.02)

        files = [
            str(p)
            for p in Path(folder).rglob("*")
            if p.suffix.lower() in SUPPORTED_EXTENSIONS
        ]
        if not files:
            return IngestResponse(success=False, message=f"Không tìm thấy file nào trong '{folder}'.")

        self._emit_progress(progress_callback, f"Tìm thấy {len(files)} file, bắt đầu ingestion...", 0.05)
        total_chunks = 0

        for file_index, file_path in enumerate(files, start=1):
            try:
                fname = os.path.basename(file_path)
                base_progress = (file_index - 1) / len(files)
                self._emit_progress(
                    progress_callback,
                    f"[{file_index}/{len(files)}] Đang đọc: {fname}",
                    base_progress,
                )
                chunks_created = self._ingest_single_file_path(
                    file_path=file_path,
                    collection=collection,
                    knowledge_base=knowledge_base,
                    owner_id=owner_id,
                    progress_callback=progress_callback,
                    file_index=file_index,
                    total_files=len(files),
                )
                total_chunks += chunks_created
            except Exception as e:
                logger.exception("Ingest error %s", file_path)
                self._emit_progress(
                    progress_callback,
                    f"[!] Lỗi: {os.path.basename(file_path)}: {e}",
                    file_index / len(files),
                )

        self._emit_progress(progress_callback, f"Xong. {total_chunks} chunks trong {clear_message}.", 1.0)
        return IngestResponse(
            success=True,
            message=f"Index thành công {len(files)} file với {total_chunks} chunks ({clear_message}).",
            documents_loaded=len(files),
            chunks_created=total_chunks,
            files=[os.path.basename(f) for f in files],
        )

    def _ingest_single_file_path(
        self,
        file_path: str,
        collection: CollectionName,
        knowledge_base: str,
        owner_id: Optional[str],
        progress_callback: Optional[Callable[[str, float], None]],
        file_index: int,
        total_files: int,
    ) -> int:
        fname = os.path.basename(file_path)
        self._store.delete_by_source_file(collection, fname, owner_id=owner_id)

        raw_text = self._extract_text(file_path)
        if not raw_text or not raw_text.strip():
            return 0

        if collection == "hr":
            validation = validate_jd_text(raw_text, fname, self._config)
            if not validation.valid:
                raise ValueError("; ".join(validation.errors))

        chunk_size = self._config.get("chunk_size", 1200)
        overlap = self._config.get("chunk_overlap", 200)
        embed_model = self._config.get("embedding_model", "nomic-embed-text")

        chunks = split_text(raw_text, chunk_size, overlap)
        if not chunks:
            return 0

        doc_type = infer_doc_type(file_path, knowledge_base)
        stem = Path(file_path).stem
        batch: List[DocumentChunk] = []

        for chunk_index, chunk_text in enumerate(chunks, start=1):
            if progress_callback and total_files > 0:
                progress = ((file_index - 1) + (chunk_index / len(chunks))) / total_files
                self._emit_progress(
                    progress_callback,
                    f"[{file_index}/{total_files}] Embedding {fname}: chunk {chunk_index}/{len(chunks)}",
                    progress,
                )
            cid = (
                f"{owner_id}_{stem}_chunk_{chunk_index - 1}"
                if owner_id
                else f"{stem}_chunk_{chunk_index - 1}"
            )
            batch.append(
                DocumentChunk(
                    id=cid,
                    text=chunk_text,
                    source_file=fname,
                    chunk_index=chunk_index - 1,
                    embedding=embed_text(self._client, embed_model, chunk_text),
                    knowledge_base=knowledge_base,
                    owner_id=owner_id,
                    doc_type=doc_type,
                )
            )

        self._store.add_chunks(batch, collection)
        return len(batch)

    def _extract_text(self, file_path: str) -> str:
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return self._extract_pdf(file_path)
        if ext == ".docx":
            return self._extract_docx(file_path)
        if ext in (".txt", ".md"):
            return Path(file_path).read_text(encoding="utf-8")
        if ext == ".json":
            return self._extract_json(file_path)
        if ext == ".xlsx":
            return self._extract_xlsx(file_path)
        if ext == ".xls":
            return self._extract_xls(file_path)
        return ""

    def _extract_pdf(self, file_path: str) -> str:
        lines = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    lines.append(text)
        return "\n\n".join(lines)

    def _extract_docx(self, file_path: str) -> str:
        doc = DocxDocument(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def _extract_json(self, file_path: str) -> str:
        raw = Path(file_path).read_text(encoding="utf-8")
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            return raw
        lines: List[str] = []
        self._flatten_json(data, "", lines)
        return "\n".join(lines)

    def _extract_xlsx(self, file_path: str) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Cần cài openpyxl. Chạy: pip install -r requirements.txt") from exc

        workbook = load_workbook(file_path, data_only=True, read_only=True)
        lines: List[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_text = self._format_excel_row(row)
                if row_text:
                    lines.append(row_text)
            lines.append("")
        workbook.close()
        return "\n".join(lines)

    def _extract_xls(self, file_path: str) -> str:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Cần cài xlrd. Chạy: pip install -r requirements.txt") from exc

        workbook = xlrd.open_workbook(file_path)
        lines: List[str] = []
        for sheet in workbook.sheets():
            lines.append(f"Sheet: {sheet.name}")
            for row_index in range(sheet.nrows):
                row_text = self._format_excel_row(sheet.row_values(row_index))
                if row_text:
                    lines.append(row_text)
            lines.append("")
        return "\n".join(lines)

    def _format_excel_row(self, row) -> str:
        values = []
        for cell in row:
            if cell is None:
                continue
            text = str(cell).strip()
            if text:
                values.append(text)
        return " | ".join(values)

    def _flatten_json(self, obj, prefix: str, lines: List[str]) -> None:
        if isinstance(obj, list):
            for item in obj:
                self._flatten_json(item, prefix, lines)
                lines.append("")
        elif isinstance(obj, dict):
            for key, value in obj.items():
                full_key = f"{prefix}.{key}" if prefix else key
                if isinstance(value, (dict, list)):
                    self._flatten_json(value, full_key, lines)
                else:
                    lines.append(f"{full_key}: {value}")
            lines.append("")
        else:
            lines.append(f"{prefix}: {obj}" if prefix else str(obj))

    def _emit_progress(
        self,
        progress_callback: Optional[Callable[[str, float], None]],
        message: str,
        progress: float,
    ) -> None:
        if progress_callback:
            progress_callback(message, max(0.0, min(progress, 1.0)))
        else:
            print(message)
